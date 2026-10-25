# -*- coding: utf-8 -*-
"""Leitura de documentos (PDF, Word, texto) e extração de campos para os laudos, com IA (Claude)."""
import base64
import json
import os
import re
import subprocess
import tempfile

ROOT = os.path.dirname(os.path.abspath(__file__))
CONFIG = os.path.join(ROOT, "dados", "config.json")
MODELO_PADRAO = "claude-sonnet-5-5"

SCHEMA_PROMPT = """Você é um assistente de Higiene Ocupacional que lê laudos e avaliações ambientais brasileiros (LTCAT, PPRA/PGR, laudos de ruído, calor, vibração e agentes químicos) e devolve APENAS um JSON com os dados que estão escritos no documento.

REGRAS
- Extraia somente o que está literalmente no documento. Não calcule, não deduza, não complete. Use null quando o dado não aparecer.
- Mantenha os números como escritos (vírgula decimal, sinais "<"). Preserve unidades.
- Se houver vários setores/funções/tabelas, devolva um item para cada um.
- "trecho" é uma citação curta do documento (máx. 160 caracteres) que sustenta o dado.
- Datas no formato DD/MM/AAAA quando possível.
- Se o documento for digitalizado e uma linha estiver ilegível, omita-a em vez de adivinhar.

FORMATO (JSON válido, sem texto fora dele):
{
 "documento": {"tipo": "", "data": "", "empresa": "", "unidade": "", "responsavel_tecnico": ""},
 "funcoes": [{"nome": "", "setor": "", "atividades": "", "ambiente": ""}],
 "ruido": [{"funcao": "", "nen_dBA": "", "dose_pct": "", "tempo_medicao_min": "", "limite_dBA": "", "trecho": ""}],
 "calor": [{"local": "", "funcao": "", "tbn": "", "tg": "", "ibutg": "", "limite": "", "data": "", "trecho": ""}],
 "quimicos": [{"setor_funcao": "", "grupo": "", "agente": "", "resultado": "", "unidade": "", "leo_acgih": "", "leo_nr15": "", "atividade": "", "periodo": "", "trecho": ""}],
 "vibracao": [{"funcao": "", "aren": "", "vdvr": "", "data": "", "trecho": ""}],
 "epi": [{"funcao": "", "epi": "", "trecho": ""}]
}
Em "quimicos", "grupo" é o título da tabela/avaliação (ex.: "Funilaria - Fumos Metálicos"), "periodo" o horário da coleta e "atividade" o que o trabalhador fazia."""


# ----------------------------------------------------------------- configuração
def carregar_config():
    try:
        with open(CONFIG, encoding="utf-8") as fh:
            return json.load(fh)
    except Exception:
        return {}


def salvar_config(cfg):
    os.makedirs(os.path.dirname(CONFIG), exist_ok=True)
    with open(CONFIG, "w", encoding="utf-8") as fh:
        json.dump(cfg, fh, ensure_ascii=False, indent=1)


def chave_api():
    return os.environ.get("ANTHROPIC_API_KEY") or carregar_config().get("anthropic_api_key") or ""


# ----------------------------------------------------------------- leitura de arquivos
def ocr_paginas(doc, paginas):
    """OCR local (Windows pt-BR) das páginas indicadas (índices a partir de 0)."""
    with tempfile.TemporaryDirectory() as tmp:
        for i in paginas:
            doc[i].get_pixmap(dpi=130).save(os.path.join(tmp, f"p{i + 1:03d}.png"))
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", os.path.join(ROOT, "ocr.ps1"), tmp],
                           capture_output=True, text=True, encoding="utf-8", errors="replace")
        return r.stdout


def ler_arquivo(nome, dados: bytes):
    """Devolve dict: texto, paginas (nº), digitalizado (bool), tipo."""
    ext = os.path.splitext(nome.lower())[1]
    if ext == ".pdf":
        import pymupdf
        doc = pymupdf.open(stream=dados, filetype="pdf")
        try:
            textos = [p.get_text() for p in doc]
            n = len(doc)
        finally:
            doc.close()
        util = sum(len(t.strip()) for t in textos)
        digit = util < 150 * n * 0.3
        sep = chr(10)
        return {"texto": "" if digit else sep.join(f"=== Página {i + 1} ==={sep}{t}" for i, t in enumerate(textos)),
                "paginas": n, "digitalizado": digit, "tipo": "pdf"}
    if ext == ".docx":
        import io
        import docx
        dd = docx.Document(io.BytesIO(dados))
        partes = [p.text for p in dd.paragraphs]
        for t in dd.tables:
            for row in t.rows:
                partes.append(" | ".join(c.text.strip() for c in row.cells))
        return {"texto": chr(10).join(partes), "paginas": 1, "digitalizado": False, "tipo": "docx"}
    return {"texto": dados.decode("utf-8", errors="replace"), "paginas": 1, "digitalizado": False, "tipo": "txt"}


def imagens_paginas(pdf_bytes, paginas, dpi=110):
    import pymupdf
    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    try:
        return [base64.b64encode(doc[i].get_pixmap(dpi=dpi).tobytes("jpeg")).decode() for i in paginas]
    finally:
        doc.close()


# ----------------------------------------------------------------- extração por IA
def _json_da_resposta(txt):
    txt = txt.strip()
    m = re.search(r"```(?:json)?\s*(.*?)```", txt, re.S)
    if m:
        txt = m.group(1)
    i, j = txt.find("{"), txt.rfind("}")
    if i < 0 or j < 0:
        raise ValueError("A resposta da IA não contém JSON.")
    return json.loads(txt[i:j + 1])


def normalizar(d):
    """Garante todas as chaves esperadas."""
    base = {"documento": {}, "funcoes": [], "ruido": [], "calor": [], "quimicos": [], "vibracao": [], "epi": []}
    for k, v in base.items():
        if k not in d or d[k] is None:
            d[k] = v
    return d


def extrair_com_ia(texto=None, imagens_b64=None, contexto="", modelo=None, cliente=None):
    """Chama a IA. `cliente` permite injetar um cliente de teste."""
    if cliente is None:
        import anthropic
        chave = chave_api()
        if not chave:
            raise RuntimeError("Chave da API não configurada (menu Configurações).")
        cliente = anthropic.Anthropic(api_key=chave)
    conteudo = []
    for b in (imagens_b64 or []):
        conteudo.append({"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": b}})
    pedido = (f"Contexto do usuário: {contexto}\n\n" if contexto else "") + "Leia o documento e devolva o JSON no formato definido."
    if texto:
        if len(texto) > 400000:
            texto = texto[:400000]
        pedido += "\n\n<documento>\n" + texto + "\n</documento>"
    conteudo.append({"type": "text", "text": pedido})
    resp = cliente.messages.create(model=modelo or carregar_config().get("modelo") or MODELO_PADRAO, max_tokens=16000,
                                   system=SCHEMA_PROMPT, messages=[{"role": "user", "content": conteudo}])
    txt = "".join(b.text for b in resp.content if getattr(b, "type", "") == "text")
    return normalizar(_json_da_resposta(txt))


# ----------------------------------------------------------------- aplicação ao laudo
def _num(s):
    try:
        return float(str(s).replace(",", ".").replace("<", "").strip())
    except Exception:
        return None


def _fmt(x, n):
    return f"{x:.{n}f}".replace(".", ",")


def aplicar_ruido(f, item):
    """item: nen_dBA, dose_pct, tempo_medicao_min. Retorna descrição do que foi aplicado."""
    import math
    D, T, NEN = _num(item.get("dose_pct")), _num(item.get("tempo_medicao_min")), _num(item.get("nen_dBA"))
    r = f.get("ruido", {})
    if D and T:
        proj = D * 480 / T
        ne = 85 + 16.61 * math.log10(D * 480 / (100 * T))
        r = {"nen": _fmt(ne, 1), "dose": _fmt(proj, 2), "fonte": r.get("fonte", "Dosimetria"), "hist": None, "acima": ne >= 80.05,
             "med": [str(item["dose_pct"]).strip(), str(item["tempo_medicao_min"]).strip(), _fmt(ne, 2)]}
    elif NEN is not None:
        dose = item.get("dose_pct") or _fmt(100 * 2 ** ((NEN - 85) / 5), 2)
        r = {"nen": _fmt(NEN, 1), "dose": str(dose), "fonte": r.get("fonte", "Dosimetria"), "hist": None, "acima": NEN >= 80.05}
    else:
        return None
    f["ruido"] = r
    return f"Ruído: NE {r['nen']} dB(A), dose {r['dose']}%"


def aplicar_quimicos(f, itens):
    """itens: lista de dicts da extração (já filtrados/editados). Monta as tabelas quantitativas."""
    grupos = {}
    for it in itens:
        if not str(it.get("agente", "")).strip():
            continue
        g = (str(it.get("grupo") or "Resultados"), str(it.get("unidade") or ""))
        grupos.setdefault(g, []).append(it)
    res = []
    for (cap, un), rs in grupos.items():
        res.append([cap, ["Agente químico", f"Resultado ({un})", f"LEO ACGIH ({un})", f"LEO NR-15 ({un})"],
                    [[x["agente"], str(x.get("resultado") or ""), str(x.get("leo_acgih") or "–"), str(x.get("leo_nr15") or "–")] for x in rs]])
    q = f.get("quim", {})
    q.update({"rows": None, "res": res})
    q.setdefault("intro", "")
    q.setdefault("concl", "")
    q.setdefault("coleta_txt", "A avaliação quantitativa dos agentes químicos foi realizada em [data], por laboratório especializado, conforme descrito no item 8.")
    f["quim"] = q
    return f"Químicos: {sum(len(r[2]) for r in res)} agentes em {len(res)} tabela(s)"


def aplicar_calor(f, itens, data_leitura=""):
    rows = []
    for it in itens:
        rows.append([str(it.get("local") or it.get("funcao") or ""), str(it.get("tbn") or ""), str(it.get("tg") or ""),
                     str(it.get("ibutg") or ""), str(it.get("limite") or "26,7"), "Abaixo do limite"])
    if not rows:
        return None
    c = f.get("calor", {})
    vals = [_num(r[3]) for r in rows if _num(r[3]) is not None]
    c.update({"rows": rows, "max": _fmt(max(vals), 2) if vals else "", "data_leitura": data_leitura or c.get("data_leitura", "")})
    f["calor"] = c
    return f"Calor: {len(rows)} leitura(s) de IBUTG"


def aplicar_atividades(f, item):
    if item.get("atividades"):
        f["atividades"] = str(item["atividades"]).strip()
    if item.get("ambiente"):
        f["ambiente"] = [p.strip() for p in re.split(r"\n\s*\n", str(item["ambiente"])) if p.strip()]
    return "Atividades/ambiente atualizados"


def registrar_origem(f, campo, doc, trecho=""):
    f.setdefault("origens", []).append({"campo": campo, "documento": doc.get("arquivo", ""), "data_documento": doc.get("data", ""),
                                        "responsavel": doc.get("responsavel_tecnico", ""), "trecho": (trecho or "")[:200]})


def ano(s):
    m = re.search(r"(19|20)\d{2}", str(s or ""))
    return int(m.group(0)) if m else None
